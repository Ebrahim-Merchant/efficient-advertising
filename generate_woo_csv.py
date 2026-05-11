"""
generate_woo_csv.py
Converts eprint_ae_product_catalogue_v8.xlsx → WooCommerce import CSV
Run: python generate_woo_csv.py
Output: woo_products_import.csv (ready to import via WP Admin > Products > Import)
"""

import sys, io, csv, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

import openpyxl

EXCEL_PATH = r'competitor_research\eprint_ae_product_catalogue_v8.xlsx'
OUTPUT_CSV  = r'woo_products_import.csv'
BRAND       = 'Efficient Advertising'
SITE_BRAND  = 'Efficient Advertising Dubai'

# ── helpers ──────────────────────────────────────────────────────────────────
def fix_brand(text):
    """Replace eprint.ae / eprint references with our brand."""
    if not text:
        return text
    text = re.sub(r'eprint\.ae', 'efficientadvt.com', text, flags=re.I)
    text = re.sub(r'\|\s*eprint\.ae', f'| {BRAND}', text, flags=re.I)
    text = re.sub(r'eprint', BRAND, text, flags=re.I)
    return text.strip()

def pipe_to_html_list(pipe_str, heading=None):
    """'A|B|C' → <ul><li>A</li>...</ul> (with optional <h4>)"""
    if not pipe_str:
        return ''
    items = [s.strip() for s in str(pipe_str).split('|') if s.strip()]
    li_html = ''.join(f'<li>✔ {i}</li>' for i in items)
    hdr = f'<h4>{heading}</h4>' if heading else ''
    return f'{hdr}<ul>{li_html}</ul>'

def comma_to_html_list(comma_str, heading=None):
    """'A,B,C' → <ul><li>A</li>...</ul>"""
    if not comma_str:
        return ''
    items = [s.strip() for s in str(comma_str).split(',') if s.strip()]
    li_html = ''.join(f'<li>{i}</li>' for i in items)
    hdr = f'<h4>{heading}</h4>' if heading else ''
    return f'{hdr}<ul>{li_html}</ul>'

def build_description(row_dict):
    """Build rich HTML product description from all Excel fields."""
    parts = []
    
    # Full description
    if row_dict.get('Full Description'):
        parts.append(f'<div class="ea-prod-desc"><p>{fix_brand(row_dict["Full Description"])}</p></div>')
    
    # Key features
    if row_dict.get('Key Features (pipe-separated)'):
        parts.append(pipe_to_html_list(row_dict['Key Features (pipe-separated)'], '🏆 Key Features'))
    
    # Materials
    if row_dict.get('Material / Substrate Options'):
        parts.append(pipe_to_html_list(row_dict['Material / Substrate Options'], '📋 Material / Substrate Options'))
    
    # Finishes
    if row_dict.get('Finish Options'):
        parts.append(pipe_to_html_list(row_dict['Finish Options'], '✨ Finish Options'))
    
    # Use cases
    if row_dict.get('Ideal Use Cases'):
        parts.append(pipe_to_html_list(row_dict['Ideal Use Cases'], '💼 Ideal Use Cases'))

    # WhatsApp CTA block
    parts.append(
        '<div class="ea-desc-cta" style="background:#f0fdf4;border:1px solid #bbf7d0;border-radius:10px;padding:20px 24px;margin-top:24px;">'
        '<p style="font-size:15px;font-weight:700;color:#166534;margin:0 0 12px;">Need a custom quote?</p>'
        '<a href="https://api.whatsapp.com/send?phone=971527966265&text=Hi%2C%20I%20need%20a%20quote%20for%20printing" '
        'target="_blank" rel="noopener" '
        'style="display:inline-flex;align-items:center;gap:8px;background:#16a34a;color:#fff;'
        'padding:12px 24px;border-radius:6px;font-size:14px;font-weight:700;text-decoration:none;">'
        '💬 WhatsApp Us for a Quote</a>'
        '</div>'
    )
    
    return '\n'.join(parts)

def make_category_hierarchy(cat, subcat):
    """Build WooCommerce category string: 'Parent > Child'"""
    if subcat and subcat.strip():
        return f'{cat.strip()} > {subcat.strip()}'
    return cat.strip() if cat else ''

def make_slug(url_slug_col, name_col):
    """Use URL slug from Excel, fallback to name-based slug."""
    if url_slug_col and str(url_slug_col).strip():
        return str(url_slug_col).strip().lower().replace(' ', '-')
    if name_col:
        return re.sub(r'[^a-z0-9-]', '-', str(name_col).lower()).strip('-')
    return ''

# ── WooCommerce CSV columns ───────────────────────────────────────────────────
WOO_HEADERS = [
    'ID', 'Type', 'SKU', 'Name', 'Published', 'Is featured?',
    'Visibility in catalog', 'Short description', 'Description',
    'Date sale price starts', 'Date sale price ends', 'Tax status', 'Tax class',
    'In stock?', 'Stock', 'Backorders allowed?', 'Sold individually?',
    'Weight (kg)', 'Length (cm)', 'Width (cm)', 'Height (cm)',
    'Allow customer reviews?', 'Purchase note', 'Sale price', 'Regular price',
    'Categories', 'Tags', 'Shipping class', 'Images', 'Download limit',
    'Download expiry days', 'Parent', 'Grouped products', 'Upsells', 'Cross-sells',
    'External URL', 'Button text', 'Position', 'Attribute 1 name', 'Attribute 1 value(s)',
    'Meta: _yoast_wpseo_title', 'Meta: _yoast_wpseo_metadesc',
    'Meta: _product_name_custom', 'slug',
]

# ── main ──────────────────────────────────────────────────────────────────────
def main():
    print(f'Loading Excel: {EXCEL_PATH}')
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
    ws = wb.active
    
    # Read headers from row 1
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    print(f'Excel columns: {headers}')
    
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    print(f'Total rows to process: {len(rows)}')
    
    # Build output CSV
    output_rows = []
    skipped = 0
    
    for i, row in enumerate(rows, start=2):
        # Map to dict
        r = {headers[j]: row[j] for j in range(len(headers)) if j < len(row)}
        
        # Skip blank rows
        if not r.get('SKU') and not r.get('Individual Product'):
            skipped += 1
            continue
        
        # Skip explicitly non-published (optional — keep all for now)
        # status = str(r.get('Product Status', 'Published')).lower()
        # if status not in ('published', 'active'): continue
        
        # Product name = Individual Product (specific variant) if exists, else Product Group
        product_name = r.get('Individual Product') or r.get('Product Group') or r.get('Sub-Category') or 'Product'
        
        # SEO Title: replace eprint branding
        seo_title = fix_brand(r.get('SEO Title (≤60 chars)', ''))
        if not seo_title:
            seo_title = f'{product_name} Dubai | {BRAND}'

        # Meta description: replace eprint branding
        meta_desc = fix_brand(r.get('Meta Description (≤155 chars)', ''))
        
        # Short description
        short_desc = fix_brand(r.get('Short Description', ''))
        
        # Categories: "Category > Sub-Category"
        cat  = r.get('Category', '')
        subcat = r.get('Sub-Category', '')
        category_str = make_category_hierarchy(cat, subcat)
        
        # Slug
        slug = make_slug(r.get('URL Slug'), product_name)
        
        # Tags from SEO Tags column (comma-separated)
        tags = r.get('SEO Tags', '') or ''
        
        # Full rich HTML description
        description = build_description(r)
        
        woo_row = {
            'ID': '',
            'Type': 'simple',
            'SKU': r.get('SKU', ''),
            'Name': str(product_name),
            'Published': '1',
            'Is featured?': '0',
            'Visibility in catalog': 'visible',
            'Short description': str(short_desc) if short_desc else '',
            'Description': description,
            'Date sale price starts': '',
            'Date sale price ends': '',
            'Tax status': 'none',
            'Tax class': '',
            'In stock?': '1',
            'Stock': '',
            'Backorders allowed?': '0',
            'Sold individually?': '0',
            'Weight (kg)': '',
            'Length (cm)': '',
            'Width (cm)': '',
            'Height (cm)': '',
            'Allow customer reviews?': '1',
            'Purchase note': '',
            'Sale price': '',
            'Regular price': '',
            'Categories': category_str,
            'Tags': str(tags),
            'Shipping class': '',
            'Images': '',
            'Download limit': '',
            'Download expiry days': '',
            'Parent': '',
            'Grouped products': '',
            'Upsells': '',
            'Cross-sells': '',
            'External URL': '',
            'Button text': '',
            'Position': '0',
            'Attribute 1 name': '',
            'Attribute 1 value(s)': '',
            'Meta: _yoast_wpseo_title': seo_title,
            'Meta: _yoast_wpseo_metadesc': meta_desc,
            'Meta: _product_name_custom': str(product_name),
            'slug': slug,
        }
        
        output_rows.append(woo_row)
    
    print(f'Products prepared: {len(output_rows)} | Skipped blank: {skipped}')
    
    # Write CSV
    with open(OUTPUT_CSV, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=WOO_HEADERS)
        writer.writeheader()
        writer.writerows(output_rows)
    
    print(f'\n✅ CSV written to: {OUTPUT_CSV}')
    print(f'   Total products: {len(output_rows)}')
    print(f'\nNext step:')
    print('  Go to: http://efficientadvt.local/wp-admin/edit.php?post_type=product')
    print('  Click: Import')
    print(f'  Upload: {OUTPUT_CSV}')
    print('  Map columns and click Import!')

if __name__ == '__main__':
    main()
