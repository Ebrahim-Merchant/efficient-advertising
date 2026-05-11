"""
EA-MASTER-EXPORT-010 — Convert product master JSON to XLSX
"""
import json
import sys
from pathlib import Path

sys.path.insert(0, r'C:\Users\pc\.python-openpyxl')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = Path(r'C:\Users\pc\Local Sites\newefficientadvertising09042026\app\public\website-cleanup-dump\logs')
IN_FILE  = BASE / 'current-product-master-export-727.json'
OUT_FILE = BASE / 'current-product-master-export-727.xlsx'

print(f'Reading {IN_FILE}...')
with open(IN_FILE, 'r', encoding='utf-8') as f:
    data = json.load(f)

rows  = data['rows']
stats = data['stats']
print(f'Rows to write: {len(rows)}')

# Column definitions: (header label, json key, column width)
COLUMNS = [
    ('Product ID',               'product_id',              10),
    ('SKU',                      'sku',                     14),
    ('Product Status',           'product_status',          14),
    ('Product Type',             'product_type',            14),
    ('Parent Product ID',        'parent_product_id',       15),
    ('Parent SKU',               'parent_sku',              14),
    ('Category Name',            'category_name',           28),
    ('Sub Category Name',        'sub_category_name',       28),
    ('Product Head',             'product_head',            38),
    ('Individual Product Name',  'individual_product_name', 42),
    ('Option / Variant Value',   'option_variant_value',    38),
    ('Regular Price',            'regular_price',           13),
    ('Sale Price',               'sale_price',              12),
    ('Short Description',        'short_description',       55),
    ('Long Description',         'long_description',        70),
    ('Meta Title',               'meta_title',              42),
    ('Meta Description',         'meta_description',        60),
    ('Featured Image ID',        'featured_image_id',       15),
    ('Featured Image File Name', 'featured_image_filename', 46),
    ('Featured Image URL',       'featured_image_url',      60),
    ('Image ALT Text',           'image_alt_text',          42),
    ('Gallery Image IDs',        'gallery_image_ids',       28),
    ('Gallery Image File Names', 'gallery_image_filenames', 60),
    ('Attribute Name',           'attribute_name',          22),
    ('Attribute Values',         'attribute_values',        70),
    ('Variation Count',          'variation_count',         14),
    ('Linked Variation SKUs',    'linked_variation_skus',   70),
    ('Merge Status',             'merge_status',            26),
    ('Notes',                    'notes',                   55),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Product Master'

# ── Header row ────────────────────────────────────────────────────────────────
header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

for col_idx, (label, _, width) in enumerate(COLUMNS, 1):
    cell = ws.cell(row=1, column=col_idx, value=label)
    cell.font  = header_font
    cell.fill  = header_fill
    cell.alignment = header_align
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.row_dimensions[1].height = 32
ws.freeze_panes = 'A2'

# ── Row fill palette ──────────────────────────────────────────────────────────
fills = {
    'variable':  PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid'),  # light green
    'variation': PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid'),  # light blue
    'draft':     PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid'),  # light orange
}

std_align = Alignment(wrap_text=False, vertical='top')
wrap_align = Alignment(wrap_text=True,  vertical='top')
WRAP_COLS  = {14, 15, 16, 17, 25, 27}  # description / SEO / attribute / variation SKU columns (1-based)

# ── Data rows ─────────────────────────────────────────────────────────────────
for row_idx, row in enumerate(rows, 2):
    ptype  = row.get('product_type', '')
    status = row.get('product_status', '')

    if status == 'draft':
        row_fill = fills['draft']
    elif ptype == 'variable':
        row_fill = fills['variable']
    elif ptype == 'variation':
        row_fill = fills['variation']
    else:
        row_fill = None

    for col_idx, (_, key, _) in enumerate(COLUMNS, 1):
        val  = row.get(key, '')
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = wrap_align if col_idx in WRAP_COLS else std_align
        if row_fill:
            cell.fill = row_fill

# ── Auto-filter on header row ─────────────────────────────────────────────────
ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}1'

# ── Stats summary sheet ───────────────────────────────────────────────────────
ws2 = wb.create_sheet('Summary')
ws2.column_dimensions['A'].width = 32
ws2.column_dimensions['B'].width = 14

summary_data = [
    ('Metric', 'Count'),
    ('Total rows exported', len(rows)),
    ('Simple products (published)', stats.get('simple', 0)),
    ('Variable products (published)', stats.get('variable', 0)),
    ('Product variations', stats.get('variation', 0)),
    ('Draft products', stats.get('draft', 0)),
    ('Products without SKU', stats.get('no_sku', 0)),
    ('Products without image', stats.get('no_image', 0)),
    ('', ''),
    ('Color Key', ''),
    ('White', 'Simple published product'),
    ('Light green', 'Variable parent product'),
    ('Light blue', 'Product variation'),
    ('Light orange', 'Draft (merged duplicate)'),
]

hdr_font = Font(bold=True)
for r_idx, (label, val) in enumerate(summary_data, 1):
    ws2.cell(row=r_idx, column=1, value=label)
    ws2.cell(row=r_idx, column=2, value=val)
    if r_idx == 1:
        ws2.cell(row=r_idx, column=1).font = hdr_font
        ws2.cell(row=r_idx, column=2).font = hdr_font

# Color swatches in summary
color_rows = {11: 'FFFFFF', 12: 'E2EFDA', 13: 'DDEBF7', 14: 'FCE4D6'}
for row_num, color in color_rows.items():
    ws2.cell(row=row_num, column=1).fill = PatternFill(start_color=color, end_color=color, fill_type='solid')

wb.save(OUT_FILE)
print(f'XLSX saved: {OUT_FILE}')
print(f'Rows: {len(rows)} | Simple: {stats["simple"]} | Variable: {stats["variable"]} | Variations: {stats["variation"]} | Draft: {stats["draft"]}')
