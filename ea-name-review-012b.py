"""
EA-NAME-REVIEW-012B: Auto-compare Product Head vs Current Website Name
REPORT ONLY — no database changes
"""
import sys, json, html
from collections import Counter

sys.path.insert(0, r'C:\Users\pc\.python-openpyxl')
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

EXCEL_PATH  = r'C:\Users\pc\OneDrive\Downloads\Edge Browser Downloads New\eprint_ae_product_catalogue_v8.xlsx'
JSON_PATH   = r'C:\Users\pc\Local Sites\newefficientadvertising09042026\app\public\website-cleanup-dump\logs\current-product-master-export-727.json'
OUTPUT_PATH = r'C:\Users\pc\Local Sites\newefficientadvertising09042026\app\public\website-cleanup-dump\logs\product-name-comparison-review.xlsx'

# ── 1. Read Excel ────────────────────────────────────────────────────────────
print('Reading Excel...')
wb_src = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
ws_src = wb_src.active
excel_all_rows = list(ws_src.iter_rows(values_only=True))
wb_src.close()

# Header: SKU=0, Category=1, Sub-Category=2, Product Group=3, Individual Product=4
excel_header = excel_all_rows[0]
print(f'Excel header: {excel_header[:5]}')
print(f'Excel data rows: {len(excel_all_rows) - 1}')

excel_data = {}   # SKU → dict
excel_skus_ordered = []  # preserve order
for row in excel_all_rows[1:]:
    sku = str(row[0]).strip() if row[0] else ''
    if not sku:
        continue
    excel_data[sku] = {
        'category':           str(row[1]).strip() if row[1] else '',
        'sub_category':       str(row[2]).strip() if row[2] else '',
        'product_group':      str(row[3]).strip() if row[3] else '',
        'individual_product': str(row[4]).strip() if row[4] else '',
    }
    excel_skus_ordered.append(sku)

print(f'Excel SKUs loaded: {len(excel_data)}')

# ── 2. Read JSON export ───────────────────────────────────────────────────────
print('Reading JSON export...')
with open(JSON_PATH, encoding='utf-8') as f:
    export = json.load(f)
json_rows = export['rows']
print(f'JSON rows: {len(json_rows)}')

# Build SKU → row index (first occurrence wins)
json_by_sku = {}
for r in json_rows:
    sku = r.get('sku', '').strip()
    if sku and sku not in json_by_sku:
        json_by_sku[sku] = r

def h(s):
    """Decode HTML entities & strip."""
    return html.unescape(str(s)).strip() if s else ''

def norm(s):
    """Normalise for comparison: strip whitespace + collapse interior spaces."""
    return ' '.join(str(s).split()) if s else ''

# ── 3. Build output rows ──────────────────────────────────────────────────────
print('Building comparison rows...')
output_rows = []
excel_skus_matched = set()

# --- Process every Excel SKU ---
for sku in excel_skus_ordered:
    ex = excel_data[sku]
    product_group_xl      = ex['product_group']
    individual_product_xl = ex['individual_product']
    category_xl           = ex['category']
    sub_category_xl       = ex['sub_category']

    if sku not in json_by_sku:
        output_rows.append({
            'sku':                       sku,
            'category':                  category_xl,
            'sub_category':              sub_category_xl,
            'product_head_excel':        product_group_xl,
            'individual_product_excel':  individual_product_xl,
            'current_website_name':      '',
            'product_type':              '',
            'parent_sku':                '',
            'parent_product_name':       '',
            'comparison_result':         'MISSING IN SITE',
            'suggested_final_name':      '',
            'notes':                     'SKU not found in current website export',
        })
        continue

    excel_skus_matched.add(sku)
    jr           = json_by_sku[sku]
    product_type = jr.get('product_type', '')
    product_status = jr.get('product_status', '')
    parent_sku   = h(jr.get('parent_sku', ''))
    product_head_json   = h(jr.get('product_head', ''))          # post_title (or parent's for variations)
    individual_name_json = h(jr.get('individual_product_name', ''))  # attribute_option for variations

    # Use Excel category if available
    category    = category_xl    or h(jr.get('category_name', ''))
    sub_category = sub_category_xl or h(jr.get('sub_category_name', ''))

    if product_type == 'variation':
        # product_head_json = parent's post_title
        # individual_name_json = attribute_option (variation label)
        current_website_name = individual_name_json
        parent_product_name  = product_head_json

        if norm(parent_product_name) == norm(product_group_xl):
            comparison     = 'CHECK VARIATION - PARENT SAME'
            suggested      = ''
        else:
            comparison     = 'CHECK VARIATION - PARENT INCORRECT'
            suggested      = product_group_xl

        notes = f'Variation of parent SKU: {parent_sku}'

    else:
        # simple (publish or draft) or variable
        current_website_name = product_head_json
        parent_product_name  = ''

        if norm(current_website_name) == norm(product_group_xl):
            comparison = 'SAME'
            suggested  = ''
        else:
            comparison = 'INCORRECT'
            suggested  = product_group_xl

        notes = 'DRAFT (merged duplicate)' if product_status == 'draft' else ''

    output_rows.append({
        'sku':                       sku,
        'category':                  category,
        'sub_category':              sub_category,
        'product_head_excel':        product_group_xl,
        'individual_product_excel':  individual_product_xl,
        'current_website_name':      current_website_name,
        'product_type':              product_type + (f' [{product_status}]' if product_status == 'draft' else ''),
        'parent_sku':                parent_sku,
        'parent_product_name':       parent_product_name,
        'comparison_result':         comparison,
        'suggested_final_name':      suggested,
        'notes':                     notes,
    })

# --- JSON SKUs not in Excel → MISSING IN EXCEL ---
for r in json_rows:
    sku = r.get('sku', '').strip()
    if not sku or sku in excel_data:
        continue
    product_type   = r.get('product_type', '')
    product_status = r.get('product_status', '')
    parent_sku     = h(r.get('parent_sku', ''))
    product_head_json    = h(r.get('product_head', ''))
    individual_name_json = h(r.get('individual_product_name', ''))

    current_website_name = individual_name_json if product_type == 'variation' else product_head_json
    parent_product_name  = product_head_json    if product_type == 'variation' else ''

    output_rows.append({
        'sku':                       sku,
        'category':                  h(r.get('category_name', '')),
        'sub_category':              h(r.get('sub_category_name', '')),
        'product_head_excel':        '',
        'individual_product_excel':  '',
        'current_website_name':      current_website_name,
        'product_type':              product_type + (f' [{product_status}]' if product_status == 'draft' else ''),
        'parent_sku':                parent_sku,
        'parent_product_name':       parent_product_name,
        'comparison_result':         'MISSING IN EXCEL',
        'suggested_final_name':      '',
        'notes':                     'SKU on website but not found in source Excel',
    })

print(f'Total output rows: {len(output_rows)}')

# ── 4. Build XLSX ─────────────────────────────────────────────────────────────
print('Writing XLSX...')
wb_out = openpyxl.Workbook()
ws = wb_out.active
ws.title = 'Name Comparison'

HEADERS = [
    'SKU',
    'Category',
    'Sub Category',
    'Product Head (Excel)',
    'Individual Product Name (Excel)',
    'Current Website Product Name',
    'Product Type',
    'Parent SKU',
    'Parent Product Name',
    'Comparison Result',
    'Suggested Final Name',
    'Notes',
]

# Header row style
HDR_FILL = PatternFill(fill_type='solid', fgColor='2F5597')
HDR_FONT = Font(bold=True, color='FFFFFF')
HDR_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)

ws.append(HEADERS)
ws.row_dimensions[1].height = 30
for cell in ws[1]:
    cell.fill  = HDR_FILL
    cell.font  = HDR_FONT
    cell.alignment = HDR_ALIGN

# Row fill colours by comparison result
FILLS = {
    'SAME':                              PatternFill(fill_type='solid', fgColor='E2EFDA'),  # green
    'INCORRECT':                         PatternFill(fill_type='solid', fgColor='FCE4D6'),  # orange/red
    'CHECK VARIATION - PARENT SAME':     PatternFill(fill_type='solid', fgColor='DDEBF7'),  # blue
    'CHECK VARIATION - PARENT INCORRECT':PatternFill(fill_type='solid', fgColor='FFF2CC'),  # yellow
    'MISSING IN SITE':                   PatternFill(fill_type='solid', fgColor='FF9999'),  # light red
    'MISSING IN EXCEL':                  PatternFill(fill_type='solid', fgColor='D9D9D9'),  # gray
}

for row_data in output_rows:
    row = [
        row_data['sku'],
        row_data['category'],
        row_data['sub_category'],
        row_data['product_head_excel'],
        row_data['individual_product_excel'],
        row_data['current_website_name'],
        row_data['product_type'],
        row_data['parent_sku'],
        row_data['parent_product_name'],
        row_data['comparison_result'],
        row_data['suggested_final_name'],
        row_data['notes'],
    ]
    ws.append(row)

    cr   = row_data['comparison_result']
    fill = FILLS.get(cr)
    if fill:
        for cell in ws[ws.max_row]:
            cell.fill = fill

# Freeze header + auto-filter
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(HEADERS))}1'

# Column widths
COL_WIDTHS = [12, 24, 24, 36, 36, 36, 18, 12, 36, 34, 36, 32]
for i, w in enumerate(COL_WIDTHS, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ── 5. Summary sheet ──────────────────────────────────────────────────────────
counts     = Counter(r['comparison_result'] for r in output_rows)
total      = len(output_rows)
same       = counts.get('SAME', 0)
incorrect  = counts.get('INCORRECT', 0)
cv_same    = counts.get('CHECK VARIATION - PARENT SAME', 0)
cv_inc     = counts.get('CHECK VARIATION - PARENT INCORRECT', 0)
miss_site  = counts.get('MISSING IN SITE', 0)
miss_excel = counts.get('MISSING IN EXCEL', 0)

ws2 = wb_out.create_sheet('Summary')

SUMMARY_ROWS = [
    ['EA-NAME-REVIEW-012B: Product Name Comparison Report', ''],
    ['Generated', '2026-05-04'],
    ['Source Excel', EXCEL_PATH],
    ['Source Export', JSON_PATH],
    ['', ''],
    ['RESULTS', ''],
    ['Metric', 'Count'],
    ['Total rows checked', total],
    ['SAME (exact match)', same],
    ['INCORRECT (name mismatch)', incorrect],
    ['CHECK VARIATION - PARENT SAME', cv_same],
    ['CHECK VARIATION - PARENT INCORRECT', cv_inc],
    ['MISSING IN SITE (in Excel, not on website)', miss_site],
    ['MISSING IN EXCEL (on website, not in Excel)', miss_excel],
    ['', ''],
    ['COLOUR KEY', ''],
    ['SAME', 'Product Head matches exactly — no action needed'],
    ['INCORRECT', 'Name mismatch — see Suggested Final Name column'],
    ['CHECK VARIATION - PARENT SAME', 'Variation; parent name matches Product Head'],
    ['CHECK VARIATION - PARENT INCORRECT', 'Variation; parent name does NOT match Product Head'],
    ['MISSING IN SITE', 'SKU in Excel but not found in website export'],
    ['MISSING IN EXCEL', 'SKU on website but not in source Excel (legacy or orphan)'],
]

for r in SUMMARY_ROWS:
    ws2.append(r)

ws2.column_dimensions['A'].width = 50
ws2.column_dimensions['B'].width = 60

ws2['A1'].font = Font(bold=True, size=14)
ws2['A6'].font = Font(bold=True)
ws2['A7'].font = Font(bold=True)
ws2['B7'].font = Font(bold=True)
ws2['A16'].font = Font(bold=True)

# Apply colour fills to colour key rows in Summary
KEY_ROW_MAP = {17: 'SAME', 18: 'INCORRECT', 19: 'CHECK VARIATION - PARENT SAME',
               20: 'CHECK VARIATION - PARENT INCORRECT', 21: 'MISSING IN SITE', 22: 'MISSING IN EXCEL'}
for row_idx, label in KEY_ROW_MAP.items():
    f = FILLS.get(label)
    if f:
        ws2.cell(row=row_idx, column=1).fill = f
        ws2.cell(row=row_idx, column=2).fill = f

wb_out.save(OUTPUT_PATH)

# ── 6. Print summary ──────────────────────────────────────────────────────────
print()
print('=' * 55)
print('EA-NAME-REVIEW-012B COMPLETE')
print('=' * 55)
print(f'Output:                          {OUTPUT_PATH}')
print(f'Total rows:                      {total}')
print(f'SAME:                            {same}')
print(f'INCORRECT:                       {incorrect}')
print(f'CHECK VARIATION - PARENT SAME:   {cv_same}')
print(f'CHECK VARIATION - PARENT INC:    {cv_inc}')
print(f'MISSING IN SITE:                 {miss_site}')
print(f'MISSING IN EXCEL:                {miss_excel}')
print('=' * 55)
