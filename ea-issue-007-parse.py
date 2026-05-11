"""
EA-MERGE-ISSUE-007 — Step 1: Excel Parse
Extracts full original_candidates + duplicate_skus for all issue groups.
Outputs JSON to website-cleanup-dump/logs/issue-007-groups.json
"""
import json
import sys
from pathlib import Path

sys.path.insert(0, r'C:\Users\pc\.python-openpyxl')
import openpyxl

SOURCE_FILE = Path(r'C:/Users/pc/Downloads/eprint_ae_product_catalogue_v8.xlsx')
OUT_DIR = Path(__file__).resolve().parent / 'website-cleanup-dump' / 'logs'
OUT_FILE = OUT_DIR / 'issue-007-groups.json'

wb = openpyxl.load_workbook(SOURCE_FILE, read_only=True, data_only=True)
ws = wb.active

headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
REMARKS_COL = 19
assert headers[REMARKS_COL] == 'Remarks', f'Expected Remarks at index 19, got: {headers[REMARKS_COL]}'

header_map = {name: idx for idx, name in enumerate(headers)}
sku_col     = header_map['SKU']
cat_col     = header_map['Category']
sub_col     = header_map['Sub-Category']
head_col    = header_map['Product Group']
name_col    = header_map['Individual Product']

rows = [row for row in ws.iter_rows(min_row=2, values_only=True)
        if row[sku_col] is not None and str(row[sku_col]).strip() != '']

# Build groups
groups = {}
for row in rows:
    key = (str(row[cat_col] or '').strip(),
           str(row[sub_col] or '').strip(),
           str(row[head_col] or '').strip())
    remark = str(row[REMARKS_COL] or '').strip().lower()
    sku = str(row[sku_col]).strip()
    name = str(row[name_col] or '').strip()
    g = groups.setdefault(key, {'originals': [], 'duplicates': []})
    if remark == 'original':
        g['originals'].append({'sku': sku, 'name': name})
    elif remark.startswith('duplicate'):
        g['duplicates'].append({'sku': sku, 'name': name, 'remark': remark.title()})

# Assign group_ids (same logic as audit script: only groups with duplicates get an id)
issue_groups = []
group_id = 0
for key, g in groups.items():
    category, sub_category, product_head = key
    if not g['duplicates']:
        continue
    group_id += 1
    n_orig = len(g['originals'])
    n_dup  = len(g['duplicates'])
    if n_orig == 0:
        issue = 'Missing Original'
    elif n_orig > 1:
        issue = 'Multiple Originals'
    else:
        issue = ''
    if issue:
        issue_groups.append({
            'group_id': group_id,
            'category': category,
            'sub_category': sub_category,
            'product_head': product_head,
            'originals': g['originals'],
            'duplicates': g['duplicates'],
            'issue_type': issue,
        })

OUT_DIR.mkdir(parents=True, exist_ok=True)
with open(OUT_FILE, 'w', encoding='utf-8') as f:
    json.dump(issue_groups, f, indent=2, ensure_ascii=False)

print(f'Issue groups found: {len(issue_groups)}')
for g in issue_groups:
    print(f"  GID {g['group_id']:2d} [{g['issue_type']:20s}] {g['product_head']} | originals: {len(g['originals'])} | dups: {len(g['duplicates'])}")
print(f'Output: {OUT_FILE}')
