import csv
import sys
from pathlib import Path

sys.path.insert(0, r'C:\Users\pc\.python-openpyxl')
import openpyxl

SOURCE_FILE = Path(r'C:/Users/pc/Downloads/eprint_ae_product_catalogue_v8.xlsx')
REPORT_DIR = Path(__file__).resolve().parent / 'website-cleanup-dump' / 'logs'
REPORT_FILE = REPORT_DIR / 'excel-variant-merge-audit.csv'

if not SOURCE_FILE.exists():
    raise FileNotFoundError(f'Excel source file not found: {SOURCE_FILE}')

wb = openpyxl.load_workbook(SOURCE_FILE, read_only=True, data_only=True)
ws = wb.active

headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))

expected_columns = ['SKU', 'Category', 'Sub-Category', 'Product Group', 'Individual Product', 'Remarks']
if len(headers) < len(expected_columns):
    raise RuntimeError(f'Unexpected Excel header count: {len(headers)}. Headers: {headers}')

# Map known column names
header_map = {name: idx for idx, name in enumerate(headers)}
for required in ['SKU', 'Category', 'Sub-Category', 'Product Group', 'Individual Product']:
    if required not in header_map:
        raise RuntimeError(f'Missing required Excel column: {required}')

# Remarks is always Column T (index 19) — confirmed from source workbook
REMARKS_COL = 19
if headers[REMARKS_COL] != 'Remarks':
    raise RuntimeError(f'Expected "Remarks" at Column T (index 19), found: {repr(headers[REMARKS_COL])}')
header_map['Remarks'] = REMARKS_COL

sku_col = header_map['SKU']
all_rows = ws.iter_rows(min_row=2, values_only=True)
# Only process rows where SKU is not blank
rows = [row for row in all_rows if row[sku_col] is not None and str(row[sku_col]).strip() != '']
total_rows = len(rows)
original_count = 0
total_duplicate_count = 0
other_remark_count = 0
missing_image_count = 0

# Build groups keyed by category/sub-category/product head
groups = {}
for row in rows:
    sku = row[header_map['SKU']]
    category = row[header_map['Category']]
    sub_category = row[header_map['Sub-Category']]
    product_head = row[header_map['Product Group']]
    product_name = row[header_map['Individual Product']]
    remark = row[header_map['Remarks']]

    key = (category, sub_category, product_head)
    group = groups.setdefault(key, {'originals': [], 'duplicates': [], 'rows': []})
    group['rows'].append(row)
    if remark and isinstance(remark, str) and remark.strip().lower() == 'original':
        group['originals'].append((sku, product_name, remark))
        original_count += 1
    elif remark and isinstance(remark, str) and remark.strip().lower().startswith('duplicate'):
        group['duplicates'].append((sku, product_name, remark))
        total_duplicate_count += 1
    else:
        # Count anything not explicitly Original or Duplicate as an other row
        group.setdefault('other_rows', []).append((sku, product_name, remark))
        other_remark_count += 1

    if any(str(cell).strip().lower() in {'missing image', 'no image', 'image missing'} for cell in row if cell is not None):
        missing_image_count += 1

if not REPORT_DIR.exists():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)

with REPORT_FILE.open('w', newline='', encoding='utf-8') as csvfile:
    writer = csv.writer(csvfile)
    writer.writerow([
        'group_id',
        'category',
        'sub_category',
        'product_head',
        'original_sku',
        'original_product_name',
        'variant_skus',
        'variant_values',
        'duplicate_count',
        'action_recommendation',
        'issues',
    ])

    group_id = 0
    groups_with_variants = 0
    groups_with_issues = 0
    expected_main_products = 0

    for key, group in groups.items():
        category, sub_category, product_head = key
        originals = group['originals']
        duplicates = group['duplicates']
        issues = []

        if duplicates:
            group_id += 1
            expected_main_products += 1

            if len(originals) == 0:
                issues.append('Missing Original')
            elif len(originals) > 1:
                issues.append('Multiple Originals')

            if not category or not sub_category or not product_head:
                issues.append('Missing grouping key')

            original_sku = originals[0][0] if originals else ''
            original_product_name = originals[0][1] if originals else ''
            variant_skus = '|'.join([sku for sku, _, _ in duplicates])
            variant_values = '|'.join([name for _, name, _ in duplicates])
            group_dup_count = len(duplicates)
            action_recommendation = 'Keep Original and convert duplicates into variants'

            if originals and duplicates:
                groups_with_variants += 1
            if issues:
                groups_with_issues += 1

            writer.writerow([
                group_id,
                category,
                sub_category,
                product_head,
                original_sku,
                original_product_name,
                variant_skus,
                variant_values,
                group_dup_count,
                action_recommendation,
                '; '.join(issues),
            ])

summary = {
    'total_valid_rows': total_rows,
    'original_rows': original_count,
    'duplicate_rows': total_duplicate_count,
    'other_remark_rows': other_remark_count,
    'expected_final_main_products': expected_main_products,
    'groups_with_variants': groups_with_variants,
    'groups_with_issues': groups_with_issues,
    'missing_image_rows': missing_image_count,
}

print('REPORT_FILE:', REPORT_FILE)
for k, v in summary.items():
    print(f'{k}: {v}')
